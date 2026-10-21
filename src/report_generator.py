"""
Export and reporting engine for candidate screening results.

Generates structured exports in CSV, Excel (.xlsx), and PDF formats,
providing in-memory bytes for Streamlit downloads and optional disk persistence.
"""

import io
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

logger = logging.getLogger(__name__)


def prepare_export_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    Format dataframe columns and list values into human-readable strings for export.

    Args:
        df: Candidate ranking DataFrame.

    Returns:
        Formatted DataFrame with strings in place of lists and clean column headers.
    """
    if df is None or df.empty:
        return pd.DataFrame()

    export_df = df.copy()

    # Format list columns into comma-separated strings
    list_cols = [
        "matched_skills",
        "missing_skills",
        "matched_preferred_skills",
        "missing_preferred_skills",
        "detected_skills",
        "required_skills",
        "preferred_skills",
    ]
    for col in list_cols:
        if col in export_df.columns:
            export_df[col] = export_df[col].apply(
                lambda val: ", ".join(val) if isinstance(val, (list, set, tuple)) else str(val)
            )

    column_mapping = {
        "rank": "Rank",
        "candidate_name": "Candidate Name",
        "strategy_used": "Strategy",
        "final_score": "Final Score (%)",
        "similarity_score": "TF-IDF Similarity (%)",
        "semantic_similarity_score": "Semantic Similarity (%)",
        "skill_match_score": "Skill Match Score (%)",
        "required_skill_match_score": "Required Skill Match (%)",
        "preferred_skill_match_score": "Preferred Skill Match (%)",
        "matched_skills": "Matched Required Skills",
        "missing_skills": "Missing Required Skills",
        "matched_preferred_skills": "Matched Preferred Skills",
        "missing_preferred_skills": "Missing Preferred Skills",
        "detected_skills": "All Detected Skills",
        "required_skills": "Required Job Skills",
        "preferred_skills": "Preferred Job Skills",
        "email": "Email",
        "phone": "Phone",
        "filename": "Filename",
        "parse_status": "Parse Status",
        "warnings": "Warnings",
    }

    # Rename existing columns
    export_df = export_df.rename(
        columns={k: v for k, v in column_mapping.items() if k in export_df.columns}
    )

    return export_df


class ReportGenerator:
    """
    Generates downloadable reports (CSV, Excel, PDF) from candidate ranking data.
    """

    def __init__(self, default_output_dir: str = "reports") -> None:
        self.output_dir = default_output_dir

    def export_csv(self, df: pd.DataFrame) -> bytes:
        """
        Export ranking results to CSV format as bytes.

        Args:
            df: Ranking DataFrame.

        Returns:
            CSV file content encoded as UTF-8 bytes.
        """
        if df is None or df.empty:
            return b""

        formatted_df = prepare_export_dataframe(df)
        csv_buffer = io.StringIO()
        formatted_df.to_csv(csv_buffer, index=False, encoding="utf-8")
        return csv_buffer.getvalue().encode("utf-8")

    def export_excel(self, df: pd.DataFrame) -> bytes:
        """
        Export ranking results to formatted Excel (.xlsx) as bytes.
        Applies bold headers, background fills, frozen header row, and column width autosizing.

        Args:
            df: Ranking DataFrame.

        Returns:
            Excel file content as bytes.
        """
        if df is None or df.empty:
            return b""

        formatted_df = prepare_export_dataframe(df)
        output = io.BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            formatted_df.to_excel(writer, index=False, sheet_name="Candidate Rankings")
            workbook = writer.book
            worksheet = writer.sheets["Candidate Rankings"]

            # Freeze top header row
            worksheet.freeze_panes = "A2"

            # Header style
            header_fill = PatternFill(
                start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
            )
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            thin_border = Border(
                left=Side(style="thin", color="D1D5DB"),
                right=Side(style="thin", color="D1D5DB"),
                top=Side(style="thin", color="D1D5DB"),
                bottom=Side(style="thin", color="D1D5DB"),
            )

            # Style header row
            for col_idx in range(1, len(formatted_df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = thin_border

            # Set row height
            worksheet.row_dimensions[1].height = 28

            # Auto-adjust column widths with bounds
            for col_idx, col_name in enumerate(formatted_df.columns, start=1):
                col_letter = get_column_letter(col_idx)
                max_len = len(str(col_name))
                for row_idx in range(2, len(formatted_df) + 2):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = min(len(val_str), 50)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output.seek(0)
        return output.getvalue()

    def export_pdf(self, df: pd.DataFrame, job_title: str = "Candidate Screening") -> bytes:
        """
        Export ranking results to a professional PDF report using ReportLab.
        Uses landscape orientation to accommodate candidate metrics and skills.

        Args:
            df: Ranking DataFrame.
            job_title: Role or screening title for report header.

        Returns:
            PDF file content as bytes.
        """
        if df is None or df.empty:
            return b""

        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="ReportTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=6,
        )

        meta_style = ParagraphStyle(
            name="ReportMeta",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#4B5563"),
            spaceAfter=10,
        )

        methodology_style = ParagraphStyle(
            name="ReportMethodology",
            parent=styles["Italic"],
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#374151"),
            spaceAfter=14,
        )

        table_header_style = ParagraphStyle(
            name="TableHeader",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.white,
            fontName="Helvetica-Bold",
            alignment=1,  # Center
        )

        table_body_style = ParagraphStyle(
            name="TableBody",
            parent=styles["Normal"],
            fontSize=7,
            leading=9,
            textColor=colors.HexColor("#111827"),
        )

        table_body_center = ParagraphStyle(
            name="TableBodyCenter",
            parent=table_body_style,
            alignment=1,
        )

        elements = []

        # Document Header
        elements.append(
            Paragraph(f"AI-Powered Candidate Screening & Ranking Report", title_style)
        )
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        elements.append(
            Paragraph(
                f"<b>Role / Batch:</b> {job_title} | <b>Generated on:</b> {timestamp} | <b>Total Candidates:</b> {len(df)}",
                meta_style,
            )
        )

        elements.append(
            Paragraph(
                "<b>Methodology Note:</b> Hybrid matching score calculated as: "
                "<i>Final Score = 0.70 × TF-IDF Cosine Similarity + 0.30 × Required Skill Match Score</i>. "
                "This report is a transparent screening decision-support aid based on lexical document analysis "
                "and does not predict job performance or replace holistic human review.",
                methodology_style,
            )
        )

        # Build Table Data
        # Column headers
        table_headers = [
            Paragraph("<b>Rank</b>", table_header_style),
            Paragraph("<b>Candidate Name</b>", table_header_style),
            Paragraph("<b>Final Score</b>", table_header_style),
            Paragraph("<b>Similarity (70%)</b>", table_header_style),
            Paragraph("<b>Skills (30%)</b>", table_header_style),
            Paragraph("<b>Matched Skills</b>", table_header_style),
            Paragraph("<b>Missing Skills</b>", table_header_style),
            Paragraph("<b>Status</b>", table_header_style),
        ]

        table_rows = [table_headers]

        for _, row in df.iterrows():
            matched = row.get("matched_skills", [])
            missing = row.get("missing_skills", [])

            matched_str = ", ".join(matched) if isinstance(matched, list) else str(matched)
            missing_str = ", ".join(missing) if isinstance(missing, list) else str(missing)

            # Limit string length gracefully for cell layout
            if len(matched_str) > 120:
                matched_str = matched_str[:117] + "..."
            if len(missing_str) > 120:
                missing_str = missing_str[:117] + "..."

            table_rows.append(
                [
                    Paragraph(str(row.get("rank", "-")), table_body_center),
                    Paragraph(str(row.get("candidate_name", "N/A")), table_body_style),
                    Paragraph(f"<b>{row.get('final_score', 0):.2f}%</b>", table_body_center),
                    Paragraph(f"{row.get('similarity_score', 0):.2f}%", table_body_center),
                    Paragraph(f"{row.get('skill_match_score', 0):.2f}%", table_body_center),
                    Paragraph(matched_str if matched_str else "None", table_body_style),
                    Paragraph(missing_str if missing_str else "None", table_body_style),
                    Paragraph(str(row.get("parse_status", "success")), table_body_center),
                ]
            )

        # Usable width in landscape letter: 11 * 72 - 72 = 720 points
        col_widths = [35, 100, 60, 65, 60, 180, 160, 60]

        table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F9FAFB")],
                    ),
                ]
            )
        )

        elements.append(table)
        doc.build(elements)

        output.seek(0)
        return output.getvalue()

    def save_report(
        self,
        content: bytes,
        filename: str,
        custom_dir: Optional[str] = None,
    ) -> Tuple[bool, str]:
        """
        Save generated report bytes to disk in designated reports folder.

        Args:
            content: File bytes.
            filename: Target filename.
            custom_dir: Optional custom folder override.

        Returns:
            Tuple of (success_boolean, file_path_or_error_message).
        """
        if not content:
            return False, "Cannot save an empty report."

        target_dir = Path(custom_dir or self.output_dir)
        try:
            target_dir.mkdir(parents=True, exist_ok=True)
            target_path = target_dir / filename
            with open(target_path, "wb") as f:
                f.write(content)
            logger.info(f"Report successfully saved to {target_path}")
            return True, str(target_path)
        except Exception as e:
            err_msg = f"Failed to save report '{filename}': {str(e)}"
            logger.error(err_msg)
            return False, err_msg
